#!/usr/bin/env python3
"""
Zefix (Switzerland) Company Lookup
----------------------------------
Company data from Zefix, the central index of the Swiss commercial
registers, run by the Federal Office of Justice. Companion to
companies_house_kyc.py (UK) and eu_company_registries.py (France, Norway).

ACCESS
    Free, but you need a username and password. Request them by email to
    zefix@bj.admin.ch (say who you are and what you'll use it for).
    The API uses HTTP Basic auth with those credentials.

    Production:  https://www.zefix.admin.ch/ZefixPublicREST/api/v1
    Test:        https://www.zefixintg.admin.ch/ZefixPublicREST/api/v1  (--env test)

WHAT YOU GET PER COMPANY
    Name, UID (CHE-xxx.xxx.xxx), CH-ID, status (active / being cancelled /
    cancelled), legal form, legal seat, canton, address, purpose, share
    capital, audit firms, head/branch offices, takeovers, former names, and
    every Swiss Official Gazette of Commerce (SOGC/SHAB) notice about it.

    NOT included as structured data: directors, board members, signatories
    or shareholders. Zefix has no people fields. People are named inside
    the text of the SOGC notices (exported in full on the Publications
    sheet) and in the official cantonal register extract (the link is in
    the cantonal_excerpt_url column).

SETUP
    pip install requests openpyxl
    Then give your credentials by ONE of:
      PowerShell:   $env:ZEFIX_USERNAME = "you";  $env:ZEFIX_PASSWORD = "secret"
      .env file:    ZEFIX_USERNAME=you  and  ZEFIX_PASSWORD=secret  on two lines
      command line: --username you --password secret

USAGE
    python zefix_switzerland.py CHE-110.088.994
    python zefix_switzerland.py CHE-110.088.994 --format xlsx
    python zefix_switzerland.py --search "microsoft"
    python zefix_switzerland.py --search "bank" --canton ZH --limit 50 --full --format xlsx
    python zefix_switzerland.py --uids-file swiss_uids.txt --format csv

    # You have company NAMES but no UIDs (e.g. banks copied from FINMA's list):
    python zefix_switzerland.py --names-file zurich_banks.xlsx --canton ZH --format xlsx
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import time
import unicodedata
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Iterable, Optional

try:
    import requests
except ImportError:  # pragma: no cover - dependency guard
    sys.exit("Missing dependency: requests.  Install it with:  pip install requests")

ENDPOINTS = {
    "prod": "https://www.zefix.admin.ch/ZefixPublicREST/api/v1",
    "test": "https://www.zefixintg.admin.ch/ZefixPublicREST/api/v1",
}

STATUS_LABELS = {"ACTIVE": "active", "BEING_CANCELLED": "being cancelled", "CANCELLED": "cancelled"}

CANTONS = {
    "AG", "AI", "AR", "BE", "BL", "BS", "FR", "GE", "GL", "GR", "JU", "LU", "NE", "NW",
    "OW", "SG", "SH", "SO", "SZ", "TG", "TI", "UR", "VD", "VS", "ZG", "ZH",
}

COMPANY_FIELDS = [
    "country", "uid", "chid", "ehraid", "company_name", "status", "legal_form",
    "legal_seat", "canton", "address", "purpose", "share_capital", "registered_on_sogc",
    "deletion_date", "audit_firms", "head_offices", "branch_offices", "has_taken_over",
    "was_taken_over_by", "former_names", "publication_count", "cantonal_excerpt_url",
    "zefix_url", "retrieved_at",
]

MATCH_FIELDS = ["name_given", "result", "matched_name", "uid", "legal_seat", "status", "candidates"]

PUBLICATION_FIELDS = [
    "uid", "company_name", "sogc_date", "sogc_id", "canton", "journal_date",
    "change_types", "message",
]


class ZefixError(Exception):
    """Anything the user can fix themselves -- printed without a traceback."""


# ---------------------------------------------------------------------------
# Credentials
# ---------------------------------------------------------------------------
def load_credentials(cli_user: Optional[str] = None, cli_password: Optional[str] = None,
                     env_file: str = ".env") -> tuple:
    """Resolve credentials from the command line, environment, or a local .env."""
    user = (cli_user or os.environ.get("ZEFIX_USERNAME") or "").strip()
    password = (cli_password or os.environ.get("ZEFIX_PASSWORD") or "").strip()

    if (not user or not password) and Path(env_file).is_file():
        for line in Path(env_file).read_text(encoding="utf-8").splitlines():
            name, sep, value = line.strip().partition("=")
            if not sep or name.startswith("#"):
                continue
            value = value.strip().strip('"').strip("'")
            if name.strip() == "ZEFIX_USERNAME" and not user:
                user = value
            elif name.strip() == "ZEFIX_PASSWORD" and not password:
                password = value

    if not user or not password:
        raise ZefixError(
            "No Zefix username/password found.\n"
            "  Zefix needs credentials -- request them by email to zefix@bj.admin.ch.\n"
            "  Then set them (PowerShell):\n"
            '    $env:ZEFIX_USERNAME = "your-username"\n'
            '    $env:ZEFIX_PASSWORD = "your-password"\n'
            "  or put ZEFIX_USERNAME=... and ZEFIX_PASSWORD=... in a .env file next to this script."
        )
    return user, password


# ---------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------
class ZefixClient:
    """Thin client: Basic auth, pacing, retries, plain-English errors."""

    MIN_INTERVAL = 0.3  # Zefix publishes no rate limit; stay polite

    def __init__(self, user: str, password: str, env: str = "prod") -> None:
        self.base = ENDPOINTS[env]
        self.session = requests.Session()
        self.session.auth = (user, password)
        self.session.headers.update({"Accept": "application/json", "User-Agent": "kyb-lookup/1.0"})
        self._last = 0.0

    def _request(self, method: str, path: str, *, body: Optional[dict] = None,
                 not_found_ok: bool = False, retries: int = 4):
        url = f"{self.base}{path}"
        delay = 2.0
        for attempt in range(retries + 1):
            wait = self.MIN_INTERVAL - (time.monotonic() - self._last)
            if wait > 0:
                time.sleep(wait)
            self._last = time.monotonic()
            try:
                resp = self.session.request(method, url, json=body, timeout=30)
            except requests.exceptions.RequestException as exc:
                if attempt == retries:
                    raise ZefixError(f"Could not reach Zefix ({exc}).") from exc
                time.sleep(delay)
                delay *= 2
                continue

            if resp.status_code == 200:
                if not resp.content:
                    return []
                try:
                    return resp.json()
                except ValueError as exc:
                    raise ZefixError(f"Zefix returned non-JSON for {path}") from exc
            if resp.status_code == 404 and not_found_ok:
                return None
            if resp.status_code == 401:
                raise ZefixError(
                    "401 Unauthorized -- Zefix rejected the username/password.\n"
                    "  * Check both were copied exactly (no spaces).\n"
                    "  * Credentials can be issued for the TEST environment only -- try --env test,\n"
                    "    or ask zefix@bj.admin.ch whether yours are enabled for production.\n"
                    "  * On a corporate network, a proxy can strip the login header; try off-network."
                )
            if resp.status_code == 403:
                raise ZefixError("403 Forbidden -- your account is not allowed to use this endpoint.")
            if resp.status_code == 429 or resp.status_code >= 500:
                if attempt == retries:
                    raise ZefixError(f"Zefix returned HTTP {resp.status_code} -- try again later.")
                time.sleep(float(resp.headers.get("Retry-After") or delay))
                delay *= 2
                continue
            if resp.status_code == 400:
                raise ZefixError(f"Zefix rejected the request (400): {resp.text[:300]}")
            raise ZefixError(f"HTTP {resp.status_code} from Zefix {path}: {resp.text[:200]}")
        raise ZefixError(f"Request to {path} failed.")

    def company_by_uid(self, uid: str) -> Optional[dict]:
        """GET /company/uid/{uid}. Zefix answers with a list; empty or 404 = not found.

        The path must use the compact form (CHE105997170): Zefix does not
        recognise the dotted display form (CHE-105.997.170) there.
        """
        data = self._request("GET", f"/company/uid/{compact_uid(uid)}", not_found_ok=True)
        if isinstance(data, list):
            return data[0] if data else None
        return data or None

    def search(self, name: str, *, canton: Optional[str] = None, active_only: bool = True,
               legal_form_id: Optional[int] = None) -> list:
        """POST /company/search. Returns short records (no address, no publications)."""
        body: dict = {"name": name, "activeOnly": active_only}
        if canton:
            body["canton"] = canton
        if legal_form_id is not None:
            body["legalFormId"] = legal_form_id
        data = self._request("POST", "/company/search", body=body, not_found_ok=True)
        return data or []


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------
def format_uid(value: str) -> str:
    """Accept CHE-110.088.994, CHE110088994 or 110088994; return CHE-110.088.994."""
    digits = re.sub(r"\D", "", str(value))
    if len(digits) != 9:
        raise ZefixError(f"'{value}' is not a Swiss UID (CHE followed by 9 digits, e.g. CHE-110.088.994).")
    return f"CHE-{digits[0:3]}.{digits[3:6]}.{digits[6:9]}"


def compact_uid(value: str) -> str:
    """Any UID form -> CHE105997170, the form Zefix accepts in lookup URLs."""
    return format_uid(value).replace("-", "").replace(".", "")


def _text(value, lang: str = "en") -> str:
    """Zefix returns some labels as {de, fr, it, en}; pick one language."""
    if isinstance(value, dict):
        return value.get(lang) or value.get("de") or next((v for v in value.values() if v), "")
    return value or ""


def _names(items) -> str:
    return "; ".join(f"{c.get('name')} ({format_uid(c['uid']) if c.get('uid') else 'no UID'})"
                     for c in items or [] if c.get("name"))


def _address(a: Optional[dict]) -> str:
    if not a:
        return ""
    street = " ".join(p for p in (a.get("street"), a.get("houseNumber")) if p)
    town = " ".join(p for p in (a.get("swissZipCode"), a.get("city")) if p)
    parts = [a.get("organisation"), f"c/o {a['careOf']}" if a.get("careOf") else None,
             street, a.get("addon"), f"PO Box {a['poBox']}" if a.get("poBox") else None, town]
    return ", ".join(p for p in parts if p)


def company_row(c: dict, lang: str = "en") -> dict:
    legal_form = c.get("legalForm") or {}
    capital = c.get("capitalNominal")
    uid = format_uid(c["uid"]) if c.get("uid") else ""
    return {
        "country": "CH",
        "uid": uid,
        "chid": c.get("chid") or "",
        "ehraid": c.get("ehraid"),
        "company_name": c.get("name"),
        "status": STATUS_LABELS.get(c.get("status"), c.get("status") or ""),
        "legal_form": _text(legal_form.get("name"), lang) or _text(legal_form.get("shortName"), lang),
        "legal_seat": c.get("legalSeat"),
        "canton": c.get("canton") or "",
        "address": _address(c.get("address")),
        "purpose": c.get("purpose") or "",
        "share_capital": f"{capital} {c.get('capitalCurrency') or ''}".strip() if capital else "",
        "registered_on_sogc": c.get("sogcDate") or "",
        "deletion_date": c.get("deletionDate") or "",
        "audit_firms": _names(c.get("auditCompanies")),
        "head_offices": _names((c.get("headOffices") or []) + (c.get("furtherHeadOffices") or [])),
        "branch_offices": _names(c.get("branchOffices")),
        "has_taken_over": _names(c.get("hasTakenOver")),
        "was_taken_over_by": _names(c.get("wasTakenOverBy")),
        "former_names": "; ".join(o.get("name") for o in c.get("oldNames") or [] if o.get("name")),
        "publication_count": len(c.get("sogcPub") or []),
        "cantonal_excerpt_url": c.get("cantonalExcerptWeb") or "",
        "zefix_url": _text(c.get("zefixDetailWeb"), lang),
        "retrieved_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


def publication_rows(c: dict) -> list:
    uid = format_uid(c["uid"]) if c.get("uid") else ""
    rows = []
    for p in sorted(c.get("sogcPub") or [], key=lambda p: p.get("sogcDate") or "", reverse=True):
        rows.append({
            "uid": uid,
            "company_name": c.get("name"),
            "sogc_date": p.get("sogcDate") or "",
            "sogc_id": p.get("sogcId") or "",
            "canton": p.get("registryOfCommerceCanton") or "",
            "journal_date": p.get("registryOfCommerceJournalDate") or "",
            "change_types": ", ".join(m.get("key") or "" for m in p.get("mutationTypes") or []),
            # Strip the HTML tags Zefix sometimes leaves in notice text.
            "message": re.sub(r"<[^>]+>", " ", p.get("message") or "").strip(),
        })
    return rows


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------
def _cell(v) -> str:
    return "" if v is None else str(v)


def export(companies: list, publications: list, formats: Iterable[str], outdir: str, stem: str,
           matches: Optional[list] = None) -> list:
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    written = []
    for fmt in formats:
        if fmt == "json":
            path = out / f"{stem}.json"
            data = [dict(c, publications=[p for p in publications if p["uid"] == c["uid"]]) for c in companies]
            if matches:
                data = {"companies": data, "matching": matches}
            path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
            written.append(path)
        elif fmt == "csv":
            for label, fields, rows in (("companies", COMPANY_FIELDS, companies),
                                        ("publications", PUBLICATION_FIELDS, publications),
                                        ("matching", MATCH_FIELDS, matches or [])):
                if not rows:
                    continue
                path = out / f"{stem}_{label}.csv"
                with path.open("w", newline="", encoding="utf-8-sig") as fh:  # BOM so Excel shows ä, é
                    w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
                    w.writeheader()
                    for row in rows:
                        w.writerow({f: _cell(row.get(f)) for f in fields})
                written.append(path)
        elif fmt == "xlsx":
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Alignment, Font
            except ImportError as exc:
                raise ZefixError("xlsx output needs openpyxl:  pip install openpyxl") from exc
            path = out / f"{stem}.xlsx"
            wb = Workbook()
            wb.remove(wb.active)
            sheets = [("Companies", COMPANY_FIELDS, companies), ("Publications", PUBLICATION_FIELDS, publications)]
            if matches:
                sheets.append(("Matching", MATCH_FIELDS, matches))
            for title, fields, rows in sheets:
                ws = wb.create_sheet(title)
                ws.append(fields)
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                for row in rows:
                    ws.append([_cell(row.get(f)) for f in fields])
                ws.freeze_panes = "A2"
                for i, f in enumerate(fields, start=1):
                    longest = max([len(f)] + [len(_cell(r.get(f))) for r in rows])
                    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(longest + 2, 12), 80)
                if title == "Publications":
                    for row in ws.iter_rows(min_row=2):
                        row[-1].alignment = Alignment(wrap_text=True, vertical="top")
            wb.save(path)
            written.append(path)
        else:
            raise ZefixError(f"Unknown format '{fmt}'. Choose from: json, csv, xlsx")
    return written


def read_uids_file(path: str) -> list:
    p = Path(path)
    if not p.is_file():
        raise ZefixError(f"UID file not found: {path}")
    uids = []
    for line in p.read_text(encoding="utf-8-sig").splitlines():
        first = line.strip().split(",")[0].strip().strip('"')
        if first and not first.startswith("#") and sum(ch.isdigit() for ch in first) == 9:
            uids.append(first)  # anything without 9 digits (e.g. a header) is skipped
    if not uids:
        raise ZefixError(f"No Swiss UIDs found in {path}")
    return uids


def read_names_file(path: str) -> list:
    """Company names from the first column of a .txt, .csv or .xlsx file.

    Made for lists copied out of FINMA's register of authorised banks, or any
    list of names: blank lines, '#' comments and a header row are skipped.
    """
    p = Path(path)
    if not p.is_file():
        raise ZefixError(f"Names file not found: {path}")
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ZefixError("Reading .xlsx needs openpyxl:  pip install openpyxl") from exc
        ws = load_workbook(p, read_only=True, data_only=True).worksheets[0]
        raw = [str(row[0]).strip() for row in ws.iter_rows(values_only=True) if row and row[0] is not None]
    else:
        with p.open(encoding="utf-8-sig", newline="") as fh:
            raw = [row[0].strip() for row in csv.reader(fh) if row]
    names = [n for n in raw if n and not n.startswith("#")]
    if names and names[0].strip().lower() in {"name", "names", "company", "company name", "firma", "bank"}:
        names = names[1:]
    seen, unique = set(), []
    for n in names:
        if n.lower() not in seen:
            seen.add(n.lower())
            unique.append(n)
    if not unique:
        raise ZefixError(f"No company names found in {path}")
    return unique


def _norm(name: str) -> str:
    """Lower-case, strip accents and punctuation: 'Zürcher Kantonalbank' -> 'zurcher kantonalbank'."""
    name = unicodedata.normalize("NFKD", name or "").encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", name.lower()).split())


def match_name(name: str, hits: list, threshold: float = 0.85) -> tuple:
    """Pick the Zefix record a given name refers to.

    Returns (hit or None, result label). Only confident matches are taken:
    an exact name match (ignoring case, accents and punctuation), the only
    hit, or one clearly-best hit above the similarity threshold. Anything
    else is reported as ambiguous for a person to resolve, rather than
    guessed.
    """
    if not hits:
        return None, "no match"
    target = _norm(name)
    exact = [h for h in hits if _norm(h.get("name")) == target]
    active_exact = [h for h in exact if h.get("status") == "ACTIVE"] or exact
    if len(active_exact) == 1:
        return active_exact[0], "exact"
    if len(hits) == 1:
        return hits[0], "only hit"
    scored = sorted(((SequenceMatcher(None, target, _norm(h.get("name"))).ratio(), h) for h in hits),
                    key=lambda t: t[0], reverse=True)
    best, runner_up = scored[0], scored[1]
    if best[0] >= threshold and best[0] - runner_up[0] >= 0.05:
        return best[1], f"close match ({best[0]:.2f})"
    return None, f"ambiguous ({len(hits)} hits)"


def resolve_names(client: "ZefixClient", names: list, *, canton: Optional[str],
                  active_only: bool) -> tuple:
    """Search Zefix for each name; return (matched UIDs, matching-report rows)."""
    uids, report = [], []
    for i, name in enumerate(names, start=1):
        print(f"[{i}/{len(names)}] searching '{name}' ...", file=sys.stderr)
        try:
            hits = client.search(name, canton=canton, active_only=active_only)
        except ZefixError as exc:
            if "401" in str(exc):
                raise
            report.append({"name_given": name, "result": f"error: {str(exc).splitlines()[0]}"})
            continue
        hit, result = match_name(name, hits)
        others = [h for h in hits if h is not hit][:5]
        row = {
            "name_given": name,
            "result": result,
            "candidates": "; ".join(f"{h.get('name')} ({h.get('legalSeat')})" for h in others),
        }
        if hit:
            row.update({
                "matched_name": hit.get("name"),
                "uid": format_uid(hit["uid"]) if hit.get("uid") else "",
                "legal_seat": hit.get("legalSeat"),
                "status": STATUS_LABELS.get(hit.get("status"), hit.get("status") or ""),
            })
            if hit.get("uid"):
                uids.append(hit["uid"])
        else:
            print(f"  ! {result} -- see the Matching sheet", file=sys.stderr)
        report.append(row)
    return uids, report


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Swiss company lookup via Zefix (commercial register index). Needs free credentials from zefix@bj.admin.ch.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "examples:\n"
            "  python zefix_switzerland.py CHE-110.088.994\n"
            "  python zefix_switzerland.py CHE-110.088.994 --format xlsx\n"
            "  python zefix_switzerland.py --search microsoft\n"
            "  python zefix_switzerland.py --search bank --canton ZH --limit 50 --full --format xlsx\n"
            "  python zefix_switzerland.py --uids-file swiss_uids.txt --format csv\n"
            "  python zefix_switzerland.py --names-file zurich_banks.xlsx --canton ZH --format xlsx\n"
        ),
    )
    p.add_argument("uids", nargs="*", help="Swiss UIDs, e.g. CHE-110.088.994 (CHE110088994 also works)")
    p.add_argument("--uids-file", metavar="PATH", help="Text/CSV file with one UID per line")
    p.add_argument("--names-file", metavar="PATH",
                   help="Company NAMES (first column of .txt/.csv/.xlsx), e.g. banks copied from FINMA's list. "
                        "Each is searched in Zefix, matched, then looked up in full.")
    p.add_argument("--search", metavar="NAME", help="Search by company name")
    p.add_argument("--canton", metavar="XX", help="Limit the search to one canton, e.g. ZH, GE, VD")
    p.add_argument("--include-inactive", action="store_true", help="Search also returns cancelled companies")
    p.add_argument("--legal-form-id", type=int, metavar="ID", help="Zefix legal form ID (e.g. 4 = GmbH)")
    p.add_argument("--limit", type=int, default=100, help="Max search hits to keep (default 100)")
    p.add_argument("--full", action="store_true", help="Fetch full details (address, publications ...) for every search hit")
    p.add_argument("--username", help="Zefix username (prefer ZEFIX_USERNAME)")
    p.add_argument("--password", help="Zefix password (prefer ZEFIX_PASSWORD)")
    p.add_argument("--env", choices=sorted(ENDPOINTS), default="prod", help="prod (default) or test")
    p.add_argument("--lang", choices=["en", "de", "fr", "it"], default="en", help="Language for legal form labels")
    p.add_argument("--format", default="json", help="Comma-separated: json, csv, xlsx (default json)")
    p.add_argument("--outdir", default="zefix_output", help="Output folder (default zefix_output)")
    p.add_argument("--stem", help="Base filename (default: auto with timestamp)")
    return p


def main(argv: Optional[list] = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        formats = [f.strip().lower() for f in args.format.split(",") if f.strip()] or ["json"]
        for f in formats:
            if f not in ("json", "csv", "xlsx"):
                raise ZefixError(f"Unknown format '{f}'. Choose from: json, csv, xlsx")
        if args.canton and args.canton.upper() not in CANTONS:
            raise ZefixError(f"'{args.canton}' is not a Swiss canton code (e.g. ZH, GE, VD, BE).")

        uids = list(args.uids)
        if args.uids_file:
            uids.extend(read_uids_file(args.uids_file))
        names = read_names_file(args.names_file) if args.names_file else []
        if not uids and not args.search and not names:
            build_parser().print_help()
            return 2

        client = ZefixClient(*load_credentials(args.username, args.password), env=args.env)
        companies, publications, matches = [], [], []

        if names:
            print(f"Matching {len(names)} names against Zefix ...", file=sys.stderr)
            found, matches = resolve_names(client, names, canton=args.canton.upper() if args.canton else None,
                                           active_only=not args.include_inactive)
            print(f"  {len(found)} of {len(names)} names matched to a company.", file=sys.stderr)
            uids.extend(u for u in found if u not in uids)

        if args.search:
            print(f"Searching Zefix for '{args.search}' ...", file=sys.stderr)
            hits = client.search(args.search, canton=args.canton.upper() if args.canton else None,
                                 active_only=not args.include_inactive, legal_form_id=args.legal_form_id)
            print(f"  {len(hits)} match(es); keeping up to {args.limit}.", file=sys.stderr)
            hits = hits[: args.limit]
            if args.full:
                uids.extend(h["uid"] for h in hits if h.get("uid"))
            else:
                companies.extend(company_row(h, args.lang) for h in hits)

        for i, uid in enumerate(uids, start=1):
            print(f"[{i}/{len(uids)}] {uid} ...", file=sys.stderr)
            try:
                full = client.company_by_uid(uid)
            except ZefixError as exc:
                if "401" in str(exc):
                    raise  # wrong credentials fail every call -- stop now
                print(f"  ! Skipped {uid}: {str(exc).splitlines()[0]}", file=sys.stderr)
                continue
            if full is None:
                print(f"  ! Not found: {uid}", file=sys.stderr)
                continue
            companies.append(company_row(full, args.lang))
            publications.extend(publication_rows(full))

        if not companies and not matches:
            print("No companies found.", file=sys.stderr)
            return 1

        stem = args.stem or f"ch_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        paths = export(companies, publications, formats, args.outdir, stem, matches)
        print(f"\n{len(companies)} companies, {len(publications)} gazette notices. Written:")
        for path in paths:
            print(f"  {path}")
        unresolved = [m for m in matches if not m.get("uid")]
        if unresolved:
            print(f"\n{len(unresolved)} name(s) not matched automatically -- check the Matching sheet, "
                  "then add their UIDs to a --uids-file.")
        if args.search and not args.full:
            print("\nSearch results are short records. Add --full for address, purpose, auditors and gazette notices.")
        return 0

    except ZefixError as exc:
        print(f"\nError: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("\nInterrupted.", file=sys.stderr)
        return 130


if __name__ == "__main__":
    sys.exit(main())
