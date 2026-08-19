"""Turn an uploaded file into plain text.

Deliberately boring. The anchoring step in `analyzer.verify_quote` matches the
model's quotes against whatever string this module returns, so the text handed
to Claude and the text used for verification must be the *same* string. Never
post-process the output of these functions differently in the two paths.
"""

from __future__ import annotations

import io
import re

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class ExtractionError(RuntimeError):
    pass


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # pypdf raises a variety of parse errors
        raise ExtractionError(f"Could not open PDF: {exc}") from exc

    if reader.is_encrypted:
        # An empty user password is common on "protected" contracts.
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ExtractionError("PDF is password protected.") from exc

    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"[page {i}]\n{text}")
    joined = "\n\n".join(pages)

    if len(joined.replace(" ", "").replace("\n", "")) < 200:
        raise ExtractionError(
            "Almost no text found. This looks like a scanned PDF -- it needs OCR "
            "before it can be analysed."
        )
    return joined


def _from_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX: {exc}") from exc

    parts = [p.text for p in document.paragraphs]

    # Contracts put commercial terms (fees, SLAs, caps) in tables far more often
    # than in prose, so dropping tables would lose the most negotiated content.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _from_plaintext(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode file as text.")


def normalise(text: str) -> str:
    """Light cleanup only.

    Collapsing runs of blank lines and stripping trailing spaces makes the text
    cheaper to send and easier to read, without moving characters around inside a
    line -- which would break quote anchoring.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract(filename: str, data: bytes) -> str:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED:
        raise ExtractionError(
            f"Unsupported file type '{suffix or filename}'. Supported: {', '.join(sorted(SUPPORTED))}"
        )

    if suffix == ".pdf":
        raw = _from_pdf(data)
    elif suffix == ".docx":
        raw = _from_docx(data)
    else:
        raw = _from_plaintext(data)

    text = normalise(raw)
    if not text:
        raise ExtractionError("File contained no readable text.")
    return text
