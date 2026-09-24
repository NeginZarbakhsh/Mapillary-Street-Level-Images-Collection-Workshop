#!/usr/bin/env python3
"""
EU Company Registry Lookup -- France and Norway
-----------------------------------------------
Companion to companies_house_kyc.py for two more official government
registers that, like UK Companies House, publish company data -- including
directors -- through a free, open REST API. Neither needs an API key or
registration.

  France  (EU)   Recherche d'entreprises -- recherche-entreprises.api.gouv.fr
                 Run by the French state (DINUM). Data from INSEE's SIRENE
                 register plus the national company register (RNE).
                 Company ID: SIREN (9 digits). A 14-digit SIRET is accepted
                 and trimmed to its SIREN.

  Norway (EEA)   Enhetsregisteret -- data.brreg.no
                 Run by the Brønnøysund Register Centre. Norway is in the
                 EEA, not the EU, but follows EU company-law rules.
                 Company ID: organisasjonsnummer (9 digits).

WHAT YOU GET PER COMPANY
    Company: name, status, legal form, incorporation/closure date, registered
    address, industry code, employee band/count.
    People:  directors/officers with role, and year (France) or full date
    (Norway) of birth. France also returns share capital and, where filed,
    turnover/net result; Norway returns bankruptcy and liquidation flags.

    Neither gives a full shareholder register or exact ownership percentages
    -- same limitation as the UK data.

SETUP
    pip install requests openpyxl          (openpyxl only needed for xlsx)

USAGE
    # One company
    python eu_company_registries.py --country fr 552032534
    python eu_company_registries.py --country no 923609016 --format xlsx

    # Search by name
    python eu_company_registries.py --country fr --search "danone"

    # Filtered listing (a sector, an area, a status)
    python eu_company_registries.py --country fr --industry 62.01Z --region 75 --status active --limit 200
    python eu_company_registries.py --country no --industry 62.010 --region 0301 --limit 200 --full

    # Your own list of IDs, one per line
    python eu_company_registries.py --country no --ids-file norway_ids.txt --format csv

    # Optional: screen directors against the UK Sanctions List (needs
    # companies_house_kyc.py in the same folder)
    python eu_company_registries.py --country fr 552032534 --sanctions
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Optional

try:
    import requests
except ImportError:  # pragma: no cover - dependency guard
    sys.exit("Missing dependency: requests.  Install it with:  pip install requests")


class RegistryError(Exception):
    """Anything the user can fix themselves -- printed without a traceback."""


COMPANY_FIELDS = [
    "country", "company_id", "company_name", "status", "legal_form", "incorporated_on",
    "closed_on", "registered_address", "industry_code", "industry_description",
    "employees", "share_capital", "flags", "source_url", "retrieved_at",
]

PEOPLE_FIELDS = [
    "country", "company_id", "company_name", "name", "role", "kind",
    "birth_date_or_year", "nationality", "entity_id", "active",
    "sanctions_hit_count", "sanctions_matches",
]


# ---------------------------------------------------------------------------
# HTTP with pacing and plain-English errors
# ---------------------------------------------------------------------------
class Http:
    """requests.Session wrapper: paces calls, retries 429/5xx, explains errors."""

    def __init__(self, min_interval: float) -> None:
        self.session = requests.Session()
        self.session.headers["User-Agent"] = "kyb-lookup/1.0"
        self.min_interval = min_interval
        self._last = 0.0

    def get(self, url: str, params: Optional[dict] = None, *, accept: str = "application/json",
            not_found_ok: bool = False, retries: int = 4) -> Optional[dict]:
        delay = 2.0
        for attempt in range(retries + 1):
            wait = self.min_interval - (time.monotonic() - self._last)
            if wait > 0:
                time.sleep(wait)
            self._last = time.monotonic()

            try:
                resp = self.session.get(url, params=params, headers={"Accept": accept}, timeout=30)
            except requests.exceptions.RequestException as exc:
                if attempt == retries:
                    raise RegistryError(f"Could not reach {url.split('/')[2]} ({exc}).") from exc
                time.sleep(delay)
                delay *= 2
                continue

            if resp.status_code == 200:
                try:
                    return resp.json()
                except ValueError as exc:
                    raise RegistryError(f"Non-JSON response from {url}") from exc
            if resp.status_code in (404, 410) and not_found_ok:
                return None  # 410 = deleted entity (Norway)
            if resp.status_code == 429 or resp.status_code >= 500:
                if attempt == retries:
                    raise RegistryError(f"HTTP {resp.status_code} from {url.split('/')[2]} -- try again later.")
                time.sleep(float(resp.headers.get("Retry-After") or delay))
                delay *= 2
                continue
            if resp.status_code == 400:
                raise RegistryError(f"The registry rejected the request (400): {resp.text[:300]}")
            raise RegistryError(f"HTTP {resp.status_code} from {url}: {resp.text[:200]}")
        raise RegistryError(f"Request to {url} failed.")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _join(*parts) -> str:
    return ", ".join(str(p).strip() for p in parts if p and str(p).strip())


def _digits(value: str) -> str:
    return re.sub(r"\D", "", str(value))


# ---------------------------------------------------------------------------
# France -- recherche-entreprises.api.gouv.fr
# ---------------------------------------------------------------------------
# Field names and limits below are taken from the API's own source code
# (github.com/annuaire-entreprises-data-gouv-fr/search-api): /search takes
# q, per_page (1-25), page (1-1000), with page*per_page capped at 10,000,
# and filters activite_principale, departement, code_postal,
# etat_administratif (A = active, C = closed), nature_juridique.
# Each result already carries its dirigeants, so no extra call per company.
class France:
    code = "fr"
    name = "France"
    id_label = "SIREN"
    BASE = "https://recherche-entreprises.api.gouv.fr"
    MAX_PER_PAGE = 25
    MAX_RESULTS = 10_000
    MIN_INTERVAL = 0.2  # the service asks for a few requests a second at most

    def __init__(self, http: Http) -> None:
        self.http = http

    def normalise_id(self, value: str) -> str:
        digits = _digits(value)
        if len(digits) == 14:  # SIRET = SIREN + 5-digit establishment number
            digits = digits[:9]
        if len(digits) != 9:
            raise RegistryError(f"'{value}' is not a French SIREN (9 digits) or SIRET (14 digits).")
        return digits

    def lookup(self, company_id: str) -> Optional[tuple]:
        siren = self.normalise_id(company_id)
        payload = self.http.get(f"{self.BASE}/search", {"q": siren, "per_page": 5})
        for result in (payload or {}).get("results") or []:
            if result.get("siren") == siren:
                return self._company(result), self._people(result)
        return None

    def search(self, *, name=None, industry=None, region=None, postcode=None,
               status=None, legal_form=None, limit: int = 100, full: bool = False) -> list:
        params: dict = {}
        if name:
            params["q"] = name
        if industry:
            params["activite_principale"] = industry
        if region:
            params["departement"] = region
        if postcode:
            params["code_postal"] = postcode
        if status:
            mapped = {"active": "A", "closed": "C", "a": "A", "c": "C"}.get(status.lower())
            if not mapped:
                raise RegistryError("France --status must be 'active' or 'closed'.")
            params["etat_administratif"] = mapped
        if legal_form:
            params["nature_juridique"] = legal_form
        if not params:
            raise RegistryError("Give a --search name or at least one filter (--industry, --region, --postcode, --status, --legal-form).")

        limit = min(limit, self.MAX_RESULTS)
        out: list = []
        page = 1
        while len(out) < limit:
            per_page = min(self.MAX_PER_PAGE, limit - len(out))
            if page * per_page > self.MAX_RESULTS:
                print(f"  ! Stopped at {len(out)}: the French API returns at most 10,000 results per query.", file=sys.stderr)
                break
            payload = self.http.get(f"{self.BASE}/search", dict(params, page=page, per_page=per_page))
            results = (payload or {}).get("results") or []
            if page == 1 and payload.get("total_results") is not None:
                print(f"  {payload['total_results']:,} companies match; retrieving up to {limit:,}.", file=sys.stderr)
            for r in results:
                out.append((self._company(r), self._people(r)))  # directors come free with each result
            if len(results) < per_page or page >= (payload.get("total_pages") or page):
                break
            page += 1
        return out[:limit]

    def _company(self, r: dict) -> dict:
        siege = r.get("siege") or {}
        immat = r.get("immatriculation") or {}
        finances = r.get("finances") or {}
        flags = []
        if finances:
            year = max(finances)
            f = finances[year] or {}
            flags.append(f"turnover {year}: {f.get('ca')}; net result {year}: {f.get('resultat_net')}")
        bodacc = r.get("bodacc") or {}
        if (bodacc.get("procedure_collective") or {}).get("statut"):
            flags.append(f"insolvency procedure: {bodacc['procedure_collective']['statut']}")
        return {
            "country": "FR",
            "company_id": r.get("siren"),
            "company_name": r.get("nom_complet") or r.get("nom_raison_sociale"),
            "status": {"A": "active", "C": "closed"}.get(r.get("etat_administratif"), r.get("etat_administratif")),
            "legal_form": r.get("nature_juridique"),
            "incorporated_on": r.get("date_creation"),
            "closed_on": r.get("date_fermeture"),
            "registered_address": siege.get("adresse") or _join(siege.get("code_postal"), siege.get("libelle_commune")),
            "industry_code": r.get("activite_principale"),
            "industry_description": r.get("section_activite_principale"),
            "employees": r.get("tranche_effectif_salarie"),
            "share_capital": _join(immat.get("capital_social"), immat.get("devise_capital")) if immat.get("capital_social") is not None else "",
            "flags": "; ".join(flags),
            "source_url": f"https://annuaire-entreprises.data.gouv.fr/entreprise/{r.get('siren')}",
            "retrieved_at": _now(),
        }

    def _people(self, r: dict) -> list:
        people = []
        for d in r.get("dirigeants") or []:
            if d.get("type_dirigeant") == "personne morale":
                people.append({
                    "name": d.get("denomination"), "role": d.get("qualite"), "kind": "entity",
                    "birth_date_or_year": "", "nationality": "", "entity_id": d.get("siren") or "",
                    "active": True,
                })
            else:
                people.append({
                    "name": _join(d.get("prenoms"), d.get("nom")).replace(",", ""),
                    "role": d.get("qualite"), "kind": "person",
                    "birth_date_or_year": d.get("date_de_naissance") or d.get("annee_de_naissance") or "",
                    "nationality": d.get("nationalite") or "", "entity_id": "", "active": True,
                })
        return people


# ---------------------------------------------------------------------------
# Norway -- data.brreg.no (Enhetsregisteret)
# ---------------------------------------------------------------------------
# Endpoints, Accept headers and JSON names below match the maintained
# python-brreg client (PyPI: brreg): GET /enheter/{orgnr} (404/410 = not
# found / deleted), GET /enheter/{orgnr}/roller, GET /enheter?... with
# 0-indexed page, results under _embedded.enheter, totals under page.
class Norway:
    code = "no"
    name = "Norway"
    id_label = "organisasjonsnummer"
    BASE = "https://data.brreg.no/enhetsregisteret/api"
    ACCEPT_ENHET = "application/vnd.brreg.enhetsregisteret.enhet.v2+json;charset=UTF-8"
    ACCEPT_ROLLE = "application/vnd.brreg.enhetsregisteret.rolle.v1+json;charset=UTF-8"
    PAGE_SIZE = 100
    MAX_RESULTS = 10_000
    MIN_INTERVAL = 0.1

    def __init__(self, http: Http) -> None:
        self.http = http

    def normalise_id(self, value: str) -> str:
        digits = _digits(value)
        if len(digits) != 9:
            raise RegistryError(f"'{value}' is not a Norwegian organisasjonsnummer (9 digits).")
        return digits

    def lookup(self, company_id: str) -> Optional[tuple]:
        orgnr = self.normalise_id(company_id)
        enhet = self.http.get(f"{self.BASE}/enheter/{orgnr}", accept=self.ACCEPT_ENHET, not_found_ok=True)
        if enhet is None:
            return None
        return self._company(enhet), self._roles(orgnr)

    def search(self, *, name=None, industry=None, region=None, postcode=None,
               status=None, legal_form=None, limit: int = 100, full: bool = False) -> list:
        params: dict = {}
        if name:
            params["navn"] = name
        if industry:
            params["naeringskode"] = industry
        if region:
            params["kommunenummer"] = region
        if postcode:
            params["forretningsadresse.postnummer"] = postcode
        if legal_form:
            params["organisasjonsform"] = legal_form
        if status:
            s = status.lower()
            if s == "active":
                params.update({"konkurs": "false", "underAvvikling": "false",
                               "underTvangsavviklingEllerTvangsopplosning": "false"})
            elif s == "bankrupt":
                params["konkurs"] = "true"
            elif s == "liquidation":
                params["underAvvikling"] = "true"
            else:
                raise RegistryError("Norway --status must be 'active', 'bankrupt' or 'liquidation'.")
        if not params:
            raise RegistryError("Give a --search name or at least one filter (--industry, --region, --postcode, --status, --legal-form).")

        limit = min(limit, self.MAX_RESULTS)
        out: list = []
        page = 0
        while len(out) < limit:
            size = min(self.PAGE_SIZE, limit - len(out))
            if (page + 1) * size > self.MAX_RESULTS:
                print(f"  ! Stopped at {len(out)}: the Norwegian API returns at most 10,000 results per query.", file=sys.stderr)
                break
            try:
                payload = self.http.get(f"{self.BASE}/enheter", dict(params, page=page, size=size), accept=self.ACCEPT_ENHET)
            except RegistryError as exc:
                if out and "400" in str(exc):
                    print(f"  ! Stopped at {len(out)} results (pagination limit).", file=sys.stderr)
                    break
                raise
            enheter = ((payload or {}).get("_embedded") or {}).get("enheter") or []
            meta = (payload or {}).get("page") or {}
            if page == 0 and meta.get("totalElements") is not None:
                print(f"  {meta['totalElements']:,} companies match; retrieving up to {limit:,}.", file=sys.stderr)
            for e in enheter:
                roles = self._roles(e["organisasjonsnummer"]) if full else []
                out.append((self._company(e), roles))
            if len(enheter) < size or page + 1 >= (meta.get("totalPages") or 0):
                break
            page += 1
        return out[:limit]

    def _company(self, e: dict) -> dict:
        addr = e.get("forretningsadresse") or e.get("postadresse") or {}
        naering = e.get("naeringskode1") or {}
        form = e.get("organisasjonsform") or {}
        if e.get("slettedato"):
            status = "deleted"
        elif e.get("konkurs"):
            status = "bankrupt"
        elif e.get("underTvangsavviklingEllerTvangsopplosning"):
            status = "forced liquidation"
        elif e.get("underAvvikling"):
            status = "in liquidation"
        else:
            status = "active"
        flags = []
        if e.get("konkursdato"):
            flags.append(f"bankruptcy date: {e['konkursdato']}")
        if e.get("sisteInnsendteAarsregnskap"):
            flags.append(f"last annual accounts filed: {e['sisteInnsendteAarsregnskap']}")
        if e.get("registrertIMvaregisteret"):
            flags.append("VAT registered")
        return {
            "country": "NO",
            "company_id": e.get("organisasjonsnummer"),
            "company_name": e.get("navn"),
            "status": status,
            "legal_form": _join(form.get("kode"), form.get("beskrivelse")),
            "incorporated_on": e.get("stiftelsesdato") or e.get("registreringsdatoEnhetsregisteret"),
            "closed_on": e.get("slettedato"),
            "registered_address": _join(*(addr.get("adresse") or []), addr.get("postnummer"), addr.get("poststed"), addr.get("land")),
            "industry_code": naering.get("kode"),
            "industry_description": naering.get("beskrivelse"),
            "employees": e.get("antallAnsatte"),
            "share_capital": "",
            "flags": "; ".join(flags),
            "source_url": f"https://virksomhet.brreg.no/nb/oppslag/enheter/{e.get('organisasjonsnummer')}",
            "retrieved_at": _now(),
        }

    def _roles(self, orgnr: str) -> list:
        payload = self.http.get(f"{self.BASE}/enheter/{orgnr}/roller", accept=self.ACCEPT_ROLLE, not_found_ok=True) or {}
        people = []
        for group in payload.get("rollegrupper") or []:
            for rolle in group.get("roller") or []:
                role = (rolle.get("type") or {}).get("beskrivelse")
                active = not rolle.get("avregistrert", False)
                person = rolle.get("person")
                enhet = rolle.get("enhet")
                if person:
                    navn = person.get("navn") or {}
                    people.append({
                        "name": _join(navn.get("fornavn"), navn.get("mellomnavn"), navn.get("etternavn")).replace(",", ""),
                        "role": role, "kind": "person",
                        "birth_date_or_year": person.get("fodselsdato") or "",
                        "nationality": "", "entity_id": "",
                        "active": active and not person.get("erDoed", False),
                    })
                elif enhet:
                    people.append({
                        "name": " ".join(enhet.get("navn") or []), "role": role, "kind": "entity",
                        "birth_date_or_year": "", "nationality": "",
                        "entity_id": enhet.get("organisasjonsnummer") or "", "active": active,
                    })
        return people


REGISTRIES = {"fr": France, "no": Norway}


# ---------------------------------------------------------------------------
# Optional sanctions screening (reuses companies_house_kyc.py if present)
# ---------------------------------------------------------------------------
def load_sanctions_screener():
    try:
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        from companies_house_kyc import load_uksl_names, screen_name  # type: ignore
    except ImportError as exc:
        raise RegistryError("--sanctions needs companies_house_kyc.py in the same folder as this script.") from exc
    try:
        names = load_uksl_names()
    except Exception as exc:  # KYBError from the other module
        raise RegistryError(f"Could not load the UK Sanctions List: {exc}") from exc
    print(f"  {len(names):,} designated names loaded (UK Sanctions List).", file=sys.stderr)
    return lambda name: screen_name(name, names)


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------
def _flatten(results: list, screen=None) -> tuple:
    companies, people = [], []
    for company, persons in results:
        companies.append(company)
        for p in persons:
            hits = screen(p["name"]) if screen and p.get("name") else []
            people.append({
                "country": company["country"], "company_id": company["company_id"],
                "company_name": company["company_name"], **p,
                "sanctions_hit_count": len(hits),
                "sanctions_matches": "; ".join(f"{h['matched_name']} ({h['score']})" for h in hits),
            })
    return companies, people


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    return str(v)


def export(results: list, formats: Iterable[str], outdir: str, stem: str, screen=None) -> list:
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    companies, people = _flatten(results, screen)
    written = []
    for fmt in formats:
        if fmt == "json":
            path = out / f"{stem}.json"
            by_company = [dict(c, people=[p for p in people if p["company_id"] == c["company_id"]]) for c in companies]
            path.write_text(json.dumps(by_company, indent=2, ensure_ascii=False), encoding="utf-8")
            written.append(path)
        elif fmt == "csv":
            for label, fields, rows in (("companies", COMPANY_FIELDS, companies), ("people", PEOPLE_FIELDS, people)):
                path = out / f"{stem}_{label}.csv"
                with path.open("w", newline="", encoding="utf-8-sig") as fh:  # BOM so Excel shows accents
                    w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
                    w.writeheader()
                    for row in rows:
                        w.writerow({f: _cell(row.get(f)) for f in fields})
                written.append(path)
        elif fmt == "xlsx":
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font
            except ImportError as exc:
                raise RegistryError("xlsx output needs openpyxl:  pip install openpyxl") from exc
            path = out / f"{stem}.xlsx"
            wb = Workbook()
            wb.remove(wb.active)
            for title, fields, rows in (("Companies", COMPANY_FIELDS, companies), ("People", PEOPLE_FIELDS, people)):
                ws = wb.create_sheet(title)
                ws.append(fields)
                for c in ws[1]:
                    c.font = Font(bold=True)
                for row in rows:
                    ws.append([_cell(row.get(f)) for f in fields])
                ws.freeze_panes = "A2"
                for i, f in enumerate(fields, start=1):
                    longest = max([len(f)] + [len(_cell(r.get(f))) for r in rows])
                    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(longest + 2, 12), 60)
            wb.save(path)
            written.append(path)
        else:
            raise RegistryError(f"Unknown format '{fmt}'. Choose from: json, csv, xlsx")
    return written


def read_ids_file(path: str) -> list:
    p = Path(path)
    if not p.is_file():
        raise RegistryError(f"IDs file not found: {path}")
    ids = []
    for line in p.read_text(encoding="utf-8-sig").splitlines():
        first = line.strip().split(",")[0].strip().strip('"')
        if first and not first.startswith("#") and any(ch.isdigit() for ch in first):
            ids.append(first)  # skips a header row like "siren" or "orgnr"
    if not ids:
        raise RegistryError(f"No IDs found in {path}")
    return ids


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Company + director lookup from the French and Norwegian official registers (free, no API key).",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "examples:\n"
            "  python eu_company_registries.py --country fr 552032534\n"
            "  python eu_company_registries.py --country no 923609016 --format xlsx\n"
            "  python eu_company_registries.py --country fr --search danone\n"
            "  python eu_company_registries.py --country fr --industry 62.01Z --region 75 --status active --limit 200\n"
            "  python eu_company_registries.py --country no --industry 62.010 --region 0301 --limit 200 --full\n"
            "  python eu_company_registries.py --country no --ids-file ids.txt --format csv\n"
        ),
    )
    p.add_argument("ids", nargs="*", help="Company IDs: SIREN/SIRET for fr, organisasjonsnummer for no")
    p.add_argument("--country", required=True, choices=sorted(REGISTRIES), help="fr = France, no = Norway")
    p.add_argument("--ids-file", metavar="PATH", help="Text/CSV file of company IDs, one per line")
    p.add_argument("--search", metavar="NAME", help="Search by company name")
    p.add_argument("--industry", metavar="CODE", help="fr: NAF code e.g. 62.01Z | no: næringskode e.g. 62.010")
    p.add_argument("--region", metavar="CODE", help="fr: département e.g. 75 | no: kommunenummer e.g. 0301 (Oslo)")
    p.add_argument("--postcode", metavar="CODE", help="Postcode of the registered address")
    p.add_argument("--status", help="fr: active|closed   no: active|bankrupt|liquidation")
    p.add_argument("--legal-form", metavar="CODE", help="fr: nature juridique e.g. 5710 (SAS) | no: e.g. AS, ASA")
    p.add_argument("--limit", type=int, default=100, help="Max companies for --search/filters (default 100, API cap 10,000)")
    p.add_argument("--full", action="store_true", help="Norway: also fetch directors for every search hit (+1 request each). France includes them already.")
    p.add_argument("--sanctions", action="store_true", help="Screen people against the UK Sanctions List (needs companies_house_kyc.py alongside)")
    p.add_argument("--format", default="json", help="Comma-separated: json, csv, xlsx (default json)")
    p.add_argument("--outdir", default="registry_output", help="Output folder (default registry_output)")
    p.add_argument("--stem", help="Base filename (default: auto with timestamp)")
    return p


def main(argv: Optional[list] = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        registry = REGISTRIES[args.country](Http(REGISTRIES[args.country].MIN_INTERVAL))
        formats = [f.strip().lower() for f in args.format.split(",") if f.strip()] or ["json"]
        for f in formats:
            if f not in ("json", "csv", "xlsx"):
                raise RegistryError(f"Unknown format '{f}'. Choose from: json, csv, xlsx")

        ids = list(args.ids)
        if args.ids_file:
            ids.extend(read_ids_file(args.ids_file))
        has_filter = any([args.search, args.industry, args.region, args.postcode, args.status, args.legal_form])
        if not ids and not has_filter:
            build_parser().print_help()
            return 2

        screen = load_sanctions_screener() if args.sanctions else None
        results: list = []

        if has_filter:
            print(f"Searching the {registry.name} register ...", file=sys.stderr)
            results.extend(registry.search(
                name=args.search, industry=args.industry, region=args.region, postcode=args.postcode,
                status=args.status, legal_form=args.legal_form, limit=args.limit, full=args.full,
            ))

        for i, cid in enumerate(ids, start=1):
            print(f"[{i}/{len(ids)}] {registry.name} {registry.id_label} {cid} ...", file=sys.stderr)
            try:
                found = registry.lookup(cid)
            except RegistryError as exc:
                print(f"  ! Skipped {cid}: {str(exc).splitlines()[0]}", file=sys.stderr)
                continue
            if found is None:
                print(f"  ! Not found (or deleted): {cid}", file=sys.stderr)
                continue
            results.append(found)

        if not results:
            print("No companies found.", file=sys.stderr)
            return 1

        stem = args.stem or f"{args.country}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        paths = export(results, formats, args.outdir, stem, screen)
        people_count = sum(len(p) for _, p in results)
        print(f"\n{len(results)} companies, {people_count} directors/officers. Written:")
        for path in paths:
            print(f"  {path}")
        return 0

    except RegistryError as exc:
        print(f"\nError: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("\nInterrupted.", file=sys.stderr)
        return 130


if __name__ == "__main__":
    sys.exit(main())
