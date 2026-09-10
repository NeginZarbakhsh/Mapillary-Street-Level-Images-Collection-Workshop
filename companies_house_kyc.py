#!/usr/bin/env python3
"""
Companies House KYC/KYB Lookup Tool
------------------------------------
Pulls company profile, officers, and PSC (beneficial ownership) data from the
UK Companies House API, optionally screens the people found against the UK
Sanctions List, and writes the result out in several formats.

SETUP
    1. Register at https://developer.company-information.service.gov.uk
    2. Create an application -> "REST API key" -> copy the key
    3. pip install requests            (openpyxl is optional, see --format xlsx)
    4. Provide the key by ONE of:
         export COMPANIES_HOUSE_API_KEY=your-key      (recommended)
         echo 'COMPANIES_HOUSE_API_KEY=your-key' > .env
         python companies_house_kyc.py 00102498 --api-key your-key

USAGE
    python companies_house_kyc.py 00102498
    python companies_house_kyc.py 00102498 09659799 --format json,csv,xlsx
    python companies_house_kyc.py --search "monzo bank"
    python companies_house_kyc.py 00102498 --format all --outdir kyb_output
    python companies_house_kyc.py 00102498 --no-sanctions

NOTE ON KEYS
    Never commit a key or paste one into a screenshot/chat. If a key has been
    exposed, delete it in the Companies House developer hub and create a new
    one -- keys cannot be un-leaked.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import time
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable, Optional
from xml.sax.saxutils import escape as xml_escape

try:
    import requests
except ImportError:  # pragma: no cover - dependency guard
    sys.exit("Missing dependency: requests.  Install it with:  pip install requests")

BASE_URL = "https://api.company-information.service.gov.uk"

# The UK Sanctions List (FCDO) has been the single official source for UK
# designations since 28 Jan 2026; it replaced the OFSI Consolidated List.
# Free, no API key. The download URL is stable (not a hashed filename).
UKSL_XML_URL = "https://sanctionslist.fcdo.gov.uk/docs/UK-Sanctions-List.xml"

SUPPORTED_FORMATS = ("json", "csv", "xlsx", "md", "txt")

# Companies House allows 600 requests per 5 minutes (2/sec sustained). Bulk runs
# set this so we pace ourselves instead of relying on 429 retries.
MIN_REQUEST_INTERVAL = 0.0
_last_request_at = 0.0


# ---------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------
class KYBError(Exception):
    """Anything the user can fix themselves -- printed without a traceback."""


# ---------------------------------------------------------------------------
# API key handling
# ---------------------------------------------------------------------------
def load_api_key(cli_key: Optional[str] = None, env_file: str = ".env") -> str:
    """Resolve the API key from --api-key, the environment, or a local .env.

    Raises KYBError with actionable instructions rather than blowing up at
    import time (the original script raised on import, so *every* run failed
    with the same message even for `--help`).
    """
    key = (cli_key or "").strip()

    if not key:
        key = (os.environ.get("COMPANIES_HOUSE_API_KEY") or "").strip()

    if not key:
        path = Path(env_file)
        if path.is_file():
            for line in path.read_text(encoding="utf-8").splitlines():
                line = line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                name, _, value = line.partition("=")
                if name.strip() == "COMPANIES_HOUSE_API_KEY":
                    key = value.strip().strip('"').strip("'")
                    break

    # Tolerate a pasted "Basic abc..." header or a trailing colon from Basic auth.
    key = re.sub(r"^(basic|bearer)\s+", "", key, flags=re.IGNORECASE).rstrip(":")

    if not key:
        raise KYBError(
            "No Companies House API key found.\n"
            "  Set one of the following and re-run:\n"
            "    export COMPANIES_HOUSE_API_KEY=your-key\n"
            "    echo 'COMPANIES_HOUSE_API_KEY=your-key' > .env\n"
            "    python companies_house_kyc.py <company_number> --api-key your-key\n"
            "  Create a key at https://developer.company-information.service.gov.uk\n"
            "  (use a 'REST API key', not a streaming key)."
        )
    return key


def make_session(api_key: str) -> "requests.Session":
    """Companies House uses HTTP Basic auth: API key as username, blank password."""
    session = requests.Session()
    session.auth = (api_key, "")
    session.headers.update({"Accept": "application/json", "User-Agent": "kyb-lookup/1.0"})
    return session


# ---------------------------------------------------------------------------
# HTTP plumbing
# ---------------------------------------------------------------------------
def _get(session, url: str, params: Optional[dict] = None, *, retries: int = 4) -> dict:
    """GET with clear errors and backoff on rate limiting / transient failures.

    Companies House allows 600 requests per 5 minutes and answers 429 when you
    exceed it -- the original code let that surface as a bare HTTPError.
    """
    global _last_request_at
    delay = 2.0
    last_error = ""

    for attempt in range(retries + 1):
        if MIN_REQUEST_INTERVAL:
            wait = MIN_REQUEST_INTERVAL - (time.monotonic() - _last_request_at)
            if wait > 0:
                time.sleep(wait)
        _last_request_at = time.monotonic()
        try:
            resp = session.get(url, params=params, timeout=30)
        except requests.exceptions.RequestException as exc:
            last_error = f"network error: {exc}"
            if attempt == retries:
                raise KYBError(f"Could not reach Companies House ({last_error}).") from exc
            time.sleep(delay)
            delay *= 2
            continue

        if resp.status_code == 200:
            try:
                return resp.json()
            except ValueError as exc:
                raise KYBError(f"Companies House returned non-JSON for {url}") from exc

        if resp.status_code == 401:
            raise KYBError(
                "401 Unauthorized -- Companies House rejected the API key.\n"
                "  * Check the key was copied in full, with no spaces.\n"
                "  * Make sure it is a REST API key, not a streaming key.\n"
                "  * A key that has been shared publicly may have been revoked;\n"
                "    create a new one in the developer hub."
            )
        if resp.status_code == 403:
            raise KYBError("403 Forbidden -- the key exists but is not allowed to call this endpoint.")
        if resp.status_code == 404:
            raise KYBError(f"404 Not Found -- no such record: {url}\n  Check the company number (8 digits, e.g. 00102498).")
        if resp.status_code == 429:
            wait = float(resp.headers.get("Retry-After") or delay)
            if attempt == retries:
                raise KYBError("429 Rate limited by Companies House (600 requests / 5 minutes). Try again shortly.")
            time.sleep(wait)
            delay *= 2
            continue
        if resp.status_code >= 500:
            last_error = f"HTTP {resp.status_code}"
            if attempt == retries:
                raise KYBError(f"Companies House is returning {last_error}. Try again later.")
            time.sleep(delay)
            delay *= 2
            continue

        raise KYBError(f"HTTP {resp.status_code} from {url}: {resp.text[:200]}")

    raise KYBError(f"Request to {url} failed: {last_error}")


def _get_all_items(session, url: str, page_size: int = 100, max_items: int = 1000) -> list:
    """Follow Companies House pagination and return every item.

    The original code read only the first page, so a company with more than 35
    officers silently lost the rest -- a real gap in a KYB check.
    """
    items: list = []
    start = 0
    while len(items) < max_items:
        payload = _get(session, url, {"items_per_page": page_size, "start_index": start})
        page = payload.get("items") or []
        items.extend(page)
        total = payload.get("total_results")
        if not page or (total is not None and len(items) >= total) or len(page) < page_size:
            break
        start += len(page)
    return items[:max_items]


# ---------------------------------------------------------------------------
# Companies House endpoints
# ---------------------------------------------------------------------------
def normalise_company_number(number: str) -> str:
    """Companies House numbers are 8 characters; plain numeric ones are zero-padded.

    People routinely type '102498' for '00102498', which used to 404.
    """
    value = str(number).strip().upper().replace(" ", "")
    if value.isdigit():
        value = value.zfill(8)
    return value


def search_company(session, query: str, items: int = 5) -> dict:
    """Search by name, for when you have a company name but not a number."""
    return _get(session, f"{BASE_URL}/search/companies", {"q": query, "items_per_page": items})


SEARCH_FIELDS = [
    "company_number", "company_name", "company_status", "company_type",
    "date_of_creation", "date_of_cessation", "sic_codes", "registered_address",
]


def advanced_search(
    session,
    *,
    name_includes: Optional[str] = None,
    name_excludes: Optional[str] = None,
    sic_codes: Optional[list] = None,
    company_status: Optional[list] = None,
    company_type: Optional[list] = None,
    location: Optional[str] = None,
    incorporated_from: Optional[str] = None,
    incorporated_to: Optional[str] = None,
    dissolved_from: Optional[str] = None,
    dissolved_to: Optional[str] = None,
    limit: int = 1000,
    page_size: int = 100,
) -> list:
    """Filtered bulk listing via /advanced-search/companies.

    There is no "list every company" endpoint -- this is the closest the API
    gets: every company matching a filter (SIC code, location, status, type,
    incorporation or dissolution date range, name fragment), paginated.

    Companies House caps how deep pagination can go, so very broad filters
    return the first N thousand rather than the true total. `hits` in the
    response tells you the real total; narrow the filter (e.g. one SIC code
    at a time, or year-by-year date ranges) to work through a large set.
    For a genuinely complete national list, use the free bulk snapshot
    instead -- see bulk_snapshot_info().
    """
    params: dict = {}
    if name_includes:
        params["company_name_includes"] = name_includes
    if name_excludes:
        params["company_name_excludes"] = name_excludes
    if sic_codes:
        params["sic_codes"] = list(sic_codes)
    if company_status:
        params["company_status"] = list(company_status)
    if company_type:
        params["company_type"] = list(company_type)
    if location:
        params["location"] = location
    if incorporated_from:
        params["incorporated_from"] = incorporated_from
    if incorporated_to:
        params["incorporated_to"] = incorporated_to
    if dissolved_from:
        params["dissolved_from"] = dissolved_from
    if dissolved_to:
        params["dissolved_to"] = dissolved_to

    if not params:
        raise KYBError(
            "Advanced search needs at least one filter -- the API will not return\n"
            "  every company on the register. Try --sic 62012, --location london,\n"
            "  --status active, --incorporated-from 2024-01-01, or --name-includes bank.\n"
            "  For the complete register, download the bulk snapshot (see --bulk-info)."
        )

    items: list = []
    start = 0
    total: Optional[int] = None

    while len(items) < limit:
        page_params = dict(params, start_index=start, size=min(page_size, limit - len(items)))
        try:
            payload = _get(session, f"{BASE_URL}/advanced-search/companies", page_params)
        except KYBError as exc:
            if items and ("416" in str(exc) or "400" in str(exc)):
                # Pagination depth cap reached -- keep what we have.
                print(f"  ! Stopped at {len(items)} results (pagination limit).", file=sys.stderr)
                break
            raise
        page = payload.get("items") or []
        if total is None:
            total = payload.get("hits")
            if total is not None:
                print(f"  {total} companies match; retrieving up to {limit}.", file=sys.stderr)
        items.extend(page)
        if not page or len(page) < page_params["size"] or (total is not None and len(items) >= total):
            break
        start += len(page)

    return items[:limit]


def search_row(item: dict) -> dict:
    """Flatten one advanced-search hit into a table row."""
    return {
        "company_number": item.get("company_number"),
        "company_name": item.get("company_name") or item.get("title"),
        "company_status": item.get("company_status"),
        "company_type": item.get("company_type") or item.get("type"),
        "date_of_creation": item.get("date_of_creation"),
        "date_of_cessation": item.get("date_of_cessation"),
        "sic_codes": ", ".join(item.get("sic_codes") or []),
        "registered_address": _format_address(item.get("registered_office_address"))
        or item.get("address_snippet", ""),
    }


def read_numbers_file(path: str) -> list:
    """Read company numbers from a text or CSV file (first column, one per line).

    Lets you run a list you already have -- e.g. rows exported from the bulk
    snapshot, or a client list from your own system.
    """
    file = Path(path)
    if not file.is_file():
        raise KYBError(f"Company number file not found: {path}")

    numbers = []
    for line in file.read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        first = line.split(",")[0].strip().strip('"').strip("'")
        if first.lower() in {"company_number", "companynumber", "number"}:
            continue  # header row
        if first:
            numbers.append(first)
    if not numbers:
        raise KYBError(f"No company numbers found in {path}")
    return numbers


def bulk_snapshot_info() -> str:
    """How to get the genuinely complete register, which no API endpoint provides."""
    return (
        "Getting EVERY UK company\n"
        "------------------------\n"
        "The REST API has no 'list all companies' endpoint. Three ways to go wide:\n"
        "\n"
        "1. Filtered bulk listing (this tool, no extra downloads):\n"
        "     python companies_house_kyc.py --advanced-search --sic 62012 --status active\n"
        "     python companies_house_kyc.py --advanced-search --location manchester --limit 2000\n"
        "   Pagination is capped, so split very large sets by SIC code or by year\n"
        "   (--incorporated-from / --incorporated-to).\n"
        "\n"
        "2. Free Company Data Product -- a monthly CSV snapshot of every live company\n"
        "   (several million rows, ~400 MB zipped, no API key needed):\n"
        "     http://download.companieshouse.gov.uk/en_output.html\n"
        "   Includes number, name, address, status, SIC codes and accounts dates.\n"
        "   It does NOT include officers or PSC detail -- those are separate\n"
        "   snapshots, and the full PSC data product is on the same download site.\n"
        "   Feed its company_number column back in with --numbers-file to enrich.\n"
        "\n"
        "3. Streaming API -- a real-time firehose of changes, for keeping a copy\n"
        "   in sync once you have loaded a snapshot. Needs a separate streaming key.\n"
        "\n"
        "Rate limit: 600 requests per 5 minutes. A full KYB check is ~4 requests, so\n"
        "roughly 150 companies per 5 minutes. This tool paces itself automatically\n"
        "for bulk runs; a few thousand companies will take hours, not minutes.\n"
    )


def get_company_profile(session, company_number: str) -> dict:
    """Core record: status, incorporation date, registered address, SIC codes, type."""
    return _get(session, f"{BASE_URL}/company/{company_number}")


def get_officers(session, company_number: str) -> list:
    """Directors and secretaries, all pages."""
    return _get_all_items(session, f"{BASE_URL}/company/{company_number}/officers")


def get_psc(session, company_number: str) -> list:
    """Persons with Significant Control -- beneficial owners (25%+ or equivalent control).

    Returns [] rather than raising when a company has no PSC register
    (the endpoint 404s for some company types, which used to abort the run).
    """
    try:
        return _get_all_items(session, f"{BASE_URL}/company/{company_number}/persons-with-significant-control")
    except KYBError as exc:
        if "404" in str(exc):
            return []
        raise


def get_filing_history(session, company_number: str, items: int = 25) -> list:
    """Recent filings -- a proxy for whether the company is actively maintained."""
    try:
        payload = _get(session, f"{BASE_URL}/company/{company_number}/filing-history", {"items_per_page": items})
        return payload.get("items") or []
    except KYBError as exc:
        if "404" in str(exc):
            return []
        raise


def _format_address(address: Optional[dict]) -> str:
    if not address:
        return ""
    parts = [
        address.get("address_line_1"),
        address.get("address_line_2"),
        address.get("locality"),
        address.get("region"),
        address.get("postal_code"),
        address.get("country"),
    ]
    return ", ".join(p for p in parts if p)


def run_kyb_check(session, company_number: str, include_filings: bool = True) -> dict:
    """Pull everything needed for a basic KYB check into one summary dict."""
    company_number = normalise_company_number(company_number)
    profile = get_company_profile(session, company_number)
    officers = get_officers(session, company_number)
    psc = get_psc(session, company_number)
    filings = get_filing_history(session, company_number) if include_filings else []

    accounts = (profile.get("accounts") or {}).get("next_due")
    confirmation = (profile.get("confirmation_statement") or {}).get("next_due")

    return {
        "retrieved_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "company_name": profile.get("company_name"),
        "company_number": profile.get("company_number") or company_number,
        "status": profile.get("company_status"),
        "status_detail": profile.get("company_status_detail"),
        "incorporated_on": profile.get("date_of_creation"),
        "dissolved_on": profile.get("date_of_cessation"),
        "company_type": profile.get("type"),
        "jurisdiction": profile.get("jurisdiction"),
        "registered_address": _format_address(profile.get("registered_office_address")),
        "registered_address_parts": profile.get("registered_office_address"),
        "sic_codes": profile.get("sic_codes") or [],
        "has_insolvency_history": bool(profile.get("has_insolvency_history")),
        "has_charges": bool(profile.get("has_charges")),
        "accounts_next_due": accounts,
        "confirmation_statement_next_due": confirmation,
        "officers": [
            {
                "name": o.get("name"),
                "role": o.get("officer_role"),
                "appointed_on": o.get("appointed_on"),
                "resigned_on": o.get("resigned_on"),
                "active": not o.get("resigned_on"),
                "nationality": o.get("nationality"),
                "occupation": o.get("occupation"),
                "date_of_birth": _dob_string(o.get("date_of_birth")),
                "country_of_residence": o.get("country_of_residence"),
                "address": _format_address(o.get("address")),
            }
            for o in officers
        ],
        "beneficial_owners": [
            {
                "name": p.get("name"),
                "kind": p.get("kind"),
                "nature_of_control": "; ".join(p.get("natures_of_control") or []),
                "notified_on": p.get("notified_on"),
                "ceased_on": p.get("ceased_on"),
                "active": not p.get("ceased_on"),
                "nationality": p.get("nationality"),
                "date_of_birth": _dob_string(p.get("date_of_birth")),
                "country_of_residence": p.get("country_of_residence"),
                "address": _format_address(p.get("address")),
            }
            for p in psc
        ],
        "recent_filings": [
            {
                "date": f.get("date"),
                "category": f.get("category"),
                "type": f.get("type"),
                "description": f.get("description"),
            }
            for f in filings
        ],
    }


def _dob_string(dob: Optional[dict]) -> str:
    """Companies House gives month/year only for privacy; format what is there."""
    if not dob:
        return ""
    year, month = dob.get("year"), dob.get("month")
    if year and month:
        return f"{int(year):04d}-{int(month):02d}"
    return str(year or "")


# ---------------------------------------------------------------------------
# UK Sanctions List screening
# ---------------------------------------------------------------------------
def load_uksl_names(local_path: Optional[str] = None, cache_dir: Optional[str] = None) -> list:
    """Download (or read) the UK Sanctions List XML and return designated names.

    Each <Name> in the real schema splits a name across up to six ordered parts
    (Name1..Name6, e.g. Name1 = given name, Name6 = surname). Collecting Name1
    and Name6 tags independently across the document -- as the previous version
    did -- yields disconnected fragments rather than whole names, so the parts
    are joined within each <Name> block here. <OrganisationName>/<FullName> are
    also picked up for schema variants that use them.

    The list is cached on disk so a network blip does not stop a KYB run.
    """
    cache = Path(cache_dir or ".") / "uk_sanctions_list.xml"
    content: Optional[bytes] = None

    if local_path:
        p = Path(local_path)
        if not p.is_file():
            raise KYBError(f"Sanctions file not found: {local_path}")
        content = p.read_bytes()
    else:
        try:
            resp = requests.get(UKSL_XML_URL, timeout=60)
            resp.raise_for_status()
            content = resp.content
            try:
                cache.write_bytes(content)
            except OSError:
                pass  # cache is a convenience, not a requirement
        except Exception as exc:
            if cache.is_file():
                print(f"  ! Could not download the sanctions list ({exc}); using cached copy.", file=sys.stderr)
                content = cache.read_bytes()
            else:
                raise KYBError(
                    f"Could not download the UK Sanctions List: {exc}\n"
                    "  Run with --no-sanctions to skip screening, or download the XML\n"
                    f"  manually from {UKSL_XML_URL} and pass --sanctions-file <path>."
                ) from exc

    try:
        root = ET.fromstring(content)
    except ET.ParseError as exc:
        raise KYBError(f"The sanctions list XML could not be parsed: {exc}") from exc

    names: list = []
    for element in root.iter():
        tag = element.tag.split("}")[-1]  # strip any XML namespace
        if tag == "Name":
            parts = []
            for part_tag in ("Name1", "Name2", "Name3", "Name4", "Name5", "Name6"):
                part = element.find(part_tag)
                if part is None:  # namespaced schema variant
                    part = next((c for c in element if c.tag.split("}")[-1] == part_tag), None)
                if part is not None and part.text and part.text.strip():
                    parts.append(part.text.strip())
            if parts:
                names.append(" ".join(parts))
        elif tag in ("OrganisationName", "FullName") and element.text and element.text.strip():
            names.append(element.text.strip())

    if not names:
        raise KYBError("Parsed the sanctions list but found no names -- the published schema may have changed.")
    return sorted(set(names))


_NAME_NOISE = re.compile(r"\b(mr|mrs|ms|miss|dr|sir|the|ltd|limited|plc|llp|inc|corp|company)\b", re.IGNORECASE)


def _normalise_name(name: str) -> str:
    name = _NAME_NOISE.sub(" ", (name or "").lower())
    name = re.sub(r"[^a-z0-9\s]", " ", name)
    return " ".join(name.split())


def screen_name(name: str, sanctions_list: Iterable[str], threshold: float = 0.85) -> list:
    """Fuzzy-match a name against the sanctions list, best score first.

    Companies House returns officers as "SURNAME, Forename", so the reversed
    form is scored too -- otherwise a genuine hit scores far below threshold.

    NOT production-grade screening: real screening needs phonetic matching,
    alias handling and transliteration (e.g. rapidfuzz plus a proper
    name-matching library). Treat hits as leads to review, not conclusions,
    and treat the absence of hits as no evidence either way.
    """
    target = _normalise_name(name)
    if not target:
        return []

    variants = {target}
    if "," in (name or ""):
        surname, _, forename = name.partition(",")
        variants.add(_normalise_name(f"{forename} {surname}"))

    hits = []
    for candidate in sanctions_list:
        cand = _normalise_name(candidate)
        if not cand:
            continue
        score = max(SequenceMatcher(None, v, cand).ratio() for v in variants)
        if score >= threshold:
            hits.append({"matched_name": candidate, "score": round(score, 3)})
    return sorted(hits, key=lambda h: h["score"], reverse=True)


def screen_kyb_result(kyb_result: dict, sanctions_list: Iterable[str], threshold: float = 0.85) -> dict:
    """Screen every officer and beneficial owner in a run_kyb_check() result."""
    sanctions_list = list(sanctions_list)
    for person in kyb_result.get("officers", []):
        person["sanctions_hits"] = screen_name(person.get("name", ""), sanctions_list, threshold)
    for person in kyb_result.get("beneficial_owners", []):
        person["sanctions_hits"] = screen_name(person.get("name", ""), sanctions_list, threshold)

    kyb_result["sanctions_screening"] = {
        "list": "UK Sanctions List (FCDO)",
        "screened_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "names_in_list": len(sanctions_list),
        "threshold": threshold,
        "total_hits": sum(
            len(p.get("sanctions_hits") or [])
            for p in kyb_result.get("officers", []) + kyb_result.get("beneficial_owners", [])
        ),
    }
    return kyb_result


# ---------------------------------------------------------------------------
# Output: tabular shaping
# ---------------------------------------------------------------------------
SUMMARY_FIELDS = [
    "company_name", "company_number", "status", "status_detail", "company_type",
    "jurisdiction", "incorporated_on", "dissolved_on", "registered_address",
    "sic_codes", "has_insolvency_history", "has_charges", "accounts_next_due",
    "confirmation_statement_next_due", "officer_count", "active_officer_count",
    "beneficial_owner_count", "sanctions_hit_count", "retrieved_at",
]

PERSON_FIELDS = [
    "company_number", "company_name", "person_type", "name", "role",
    "nature_of_control", "kind", "appointed_on", "notified_on", "resigned_on",
    "ceased_on", "active", "nationality", "date_of_birth", "country_of_residence",
    "address", "sanctions_hit_count", "sanctions_matches",
]

FILING_FIELDS = ["company_number", "date", "category", "type", "description"]


def _hits_text(person: dict) -> str:
    return "; ".join(f"{h['matched_name']} ({h['score']})" for h in person.get("sanctions_hits") or [])


def summary_row(result: dict) -> dict:
    people = result.get("officers", []) + result.get("beneficial_owners", [])
    return {
        "company_name": result.get("company_name"),
        "company_number": result.get("company_number"),
        "status": result.get("status"),
        "status_detail": result.get("status_detail"),
        "company_type": result.get("company_type"),
        "jurisdiction": result.get("jurisdiction"),
        "incorporated_on": result.get("incorporated_on"),
        "dissolved_on": result.get("dissolved_on"),
        "registered_address": result.get("registered_address"),
        "sic_codes": ", ".join(result.get("sic_codes") or []),
        "has_insolvency_history": result.get("has_insolvency_history"),
        "has_charges": result.get("has_charges"),
        "accounts_next_due": result.get("accounts_next_due"),
        "confirmation_statement_next_due": result.get("confirmation_statement_next_due"),
        "officer_count": len(result.get("officers", [])),
        "active_officer_count": sum(1 for o in result.get("officers", []) if o.get("active")),
        "beneficial_owner_count": len(result.get("beneficial_owners", [])),
        "sanctions_hit_count": sum(len(p.get("sanctions_hits") or []) for p in people),
        "retrieved_at": result.get("retrieved_at"),
    }


def people_rows(result: dict) -> list:
    rows = []
    for person_type, key in (("officer", "officers"), ("beneficial_owner", "beneficial_owners")):
        for person in result.get(key, []):
            row = {field: "" for field in PERSON_FIELDS}
            row.update(
                {
                    "company_number": result.get("company_number"),
                    "company_name": result.get("company_name"),
                    "person_type": person_type,
                    "sanctions_hit_count": len(person.get("sanctions_hits") or []),
                    "sanctions_matches": _hits_text(person),
                }
            )
            for field in PERSON_FIELDS:
                if field in person:
                    row[field] = person.get(field)
            rows.append(row)
    return rows


def filing_rows(result: dict) -> list:
    return [
        {"company_number": result.get("company_number"), **filing}
        for filing in result.get("recent_filings", [])
    ]


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


# ---------------------------------------------------------------------------
# Output: writers
# ---------------------------------------------------------------------------
def write_json(results: list, outdir: Path, stem: str) -> Path:
    path = outdir / f"{stem}.json"
    payload = results[0] if len(results) == 1 else results
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def write_csv(results: list, outdir: Path, stem: str) -> list:
    """Three CSVs: companies, people, filings. Excel-safe UTF-8 (BOM)."""
    written = []
    tables = [
        ("companies", SUMMARY_FIELDS, [summary_row(r) for r in results]),
        ("people", PERSON_FIELDS, [row for r in results for row in people_rows(r)]),
        ("filings", FILING_FIELDS, [row for r in results for row in filing_rows(r)]),
    ]
    for name, fields, rows in tables:
        if not rows:
            continue
        path = outdir / f"{stem}_{name}.csv"
        with path.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({f: _stringify(row.get(f)) for f in fields})
        written.append(path)
    return written


def _write_xlsx_stdlib(path: Path, sheets: list) -> None:
    """Minimal .xlsx writer using only the standard library.

    sheets: [(sheet_name, [header,...], [[cell,...],...])]
    Uses inline strings, so no shared-string table is needed. Keeps xlsx export
    working when openpyxl is not installed (e.g. a bare Colab runtime).
    """
    import zipfile

    def col_ref(idx: int) -> str:
        ref = ""
        idx += 1
        while idx:
            idx, rem = divmod(idx - 1, 26)
            ref = chr(65 + rem) + ref
        return ref

    def sheet_xml(header: list, rows: list) -> str:
        out = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>',
        ]
        for r_idx, row in enumerate([header] + rows, start=1):
            out.append(f'<row r="{r_idx}">')
            for c_idx, value in enumerate(row):
                text = _stringify(value)
                if text == "":
                    continue
                out.append(
                    f'<c r="{col_ref(c_idx)}{r_idx}" t="inlineStr"><is><t xml:space="preserve">'
                    f"{xml_escape(text)}</t></is></c>"
                )
            out.append("</row>")
        out.append("</sheetData></worksheet>")
        return "".join(out)

    sheet_entries = "".join(
        f'<sheet name="{xml_escape(name)}" sheetId="{i}" r:id="rId{i}"/>'
        for i, (name, _, _) in enumerate(sheets, start=1)
    )
    rel_entries = "".join(
        f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>'
        for i in range(1, len(sheets) + 1)
    )
    overrides = "".join(
        f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for i in range(1, len(sheets) + 1)
    )

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            f"{overrides}</Types>",
        )
        zf.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>",
        )
        zf.writestr(
            "xl/workbook.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f"<sheets>{sheet_entries}</sheets></workbook>",
        )
        zf.writestr(
            "xl/_rels/workbook.xml.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f"{rel_entries}</Relationships>",
        )
        for i, (_, header, rows) in enumerate(sheets, start=1):
            zf.writestr(f"xl/worksheets/sheet{i}.xml", sheet_xml(header, rows))


def write_xlsx(results: list, outdir: Path, stem: str) -> Path:
    """One workbook, three sheets. Uses openpyxl when available, else stdlib."""
    path = outdir / f"{stem}.xlsx"
    sheets = [
        ("Companies", SUMMARY_FIELDS, [[summary_row(r).get(f) for f in SUMMARY_FIELDS] for r in results]),
        ("People", PERSON_FIELDS, [[row.get(f) for f in PERSON_FIELDS] for r in results for row in people_rows(r)]),
        ("Filings", FILING_FIELDS, [[row.get(f) for f in FILING_FIELDS] for r in results for row in filing_rows(r)]),
    ]

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        _write_xlsx_stdlib(path, sheets)
        return path

    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_stringify(v) for v in row])
        ws.freeze_panes = "A2"
        for idx, column in enumerate(header, start=1):
            longest = max([len(column)] + [len(_stringify(r[idx - 1])) for r in rows] or [0])
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(longest + 2, 12), 60)
    wb.save(path)
    return path


def _markdown_table(header: list, rows: list) -> str:
    if not rows:
        return "_None recorded._\n"
    out = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
    for row in rows:
        out.append("| " + " | ".join(_stringify(c).replace("|", "\\|") or "-" for c in row) + " |")
    return "\n".join(out) + "\n"


def write_markdown(results: list, outdir: Path, stem: str) -> Path:
    path = outdir / f"{stem}.md"
    chunks = ["# KYB report", ""]
    for result in results:
        summary = summary_row(result)
        chunks.append(f"## {summary['company_name']} ({summary['company_number']})")
        chunks.append("")
        chunks.append(
            _markdown_table(
                ["Field", "Value"],
                [[k.replace("_", " ").title(), summary[k]] for k in SUMMARY_FIELDS if k in summary],
            )
        )
        chunks.append("### Officers\n")
        chunks.append(
            _markdown_table(
                ["Name", "Role", "Appointed", "Active", "Nationality", "Sanctions hits"],
                [
                    [o.get("name"), o.get("role"), o.get("appointed_on"), o.get("active"), o.get("nationality"), _hits_text(o) or "none"]
                    for o in result.get("officers", [])
                ],
            )
        )
        chunks.append("### Beneficial owners (PSC)\n")
        chunks.append(
            _markdown_table(
                ["Name", "Kind", "Nature of control", "Notified", "Active", "Sanctions hits"],
                [
                    [p.get("name"), p.get("kind"), p.get("nature_of_control"), p.get("notified_on"), p.get("active"), _hits_text(p) or "none"]
                    for p in result.get("beneficial_owners", [])
                ],
            )
        )
        if result.get("recent_filings"):
            chunks.append("### Recent filings\n")
            chunks.append(
                _markdown_table(
                    ["Date", "Category", "Description"],
                    [[f.get("date"), f.get("category"), f.get("description")] for f in result["recent_filings"][:15]],
                )
            )
        chunks.append("")
    path.write_text("\n".join(chunks), encoding="utf-8")
    return path


def write_text(results: list, outdir: Path, stem: str) -> Path:
    path = outdir / f"{stem}.txt"
    lines = []
    for result in results:
        summary = summary_row(result)
        lines.append("=" * 70)
        lines.append(f"{summary['company_name']}  ({summary['company_number']})")
        lines.append("=" * 70)
        for field in SUMMARY_FIELDS:
            lines.append(f"  {field.replace('_', ' ').title():<34}{_stringify(summary.get(field))}")
        for label, key in (("OFFICERS", "officers"), ("BENEFICIAL OWNERS (PSC)", "beneficial_owners")):
            lines.append("")
            lines.append(f"  {label}")
            people = result.get(key, [])
            if not people:
                lines.append("    (none recorded)")
            for person in people:
                detail = person.get("role") or person.get("nature_of_control") or person.get("kind") or ""
                lines.append(f"    - {person.get('name')}  [{detail}]")
                hits = _hits_text(person)
                if hits:
                    lines.append(f"        SANCTIONS HIT: {hits}")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def export(results: list, formats: Iterable[str], outdir: str = "kyb_output", stem: Optional[str] = None) -> list:
    """Write the results in each requested format; returns the paths written."""
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)

    if stem is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        stem = f"kyb_{results[0].get('company_number')}_{stamp}" if len(results) == 1 else f"kyb_batch_{stamp}"

    written: list = []
    for fmt in formats:
        fmt = fmt.strip().lower()
        if fmt == "json":
            written.append(write_json(results, out, stem))
        elif fmt == "csv":
            written.extend(write_csv(results, out, stem))
        elif fmt == "xlsx":
            written.append(write_xlsx(results, out, stem))
        elif fmt == "md":
            written.append(write_markdown(results, out, stem))
        elif fmt == "txt":
            written.append(write_text(results, out, stem))
        else:
            raise KYBError(f"Unknown format '{fmt}'. Choose from: {', '.join(SUPPORTED_FORMATS)}, all")
    return written


def export_search(items: list, formats: Iterable[str], outdir: str = "kyb_output", stem: Optional[str] = None) -> list:
    """Write an advanced-search result set (a flat company list) to disk.

    Same formats as export(), but one table rather than the full KYB structure --
    this is a directory of companies, not a per-company dossier.
    """
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    stem = stem or f"companies_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    rows = [search_row(item) for item in items]
    written: list = []

    for fmt in formats:
        fmt = fmt.strip().lower()
        if fmt == "json":
            path = out / f"{stem}.json"
            path.write_text(json.dumps(rows, indent=2, ensure_ascii=False), encoding="utf-8")
        elif fmt == "csv":
            path = out / f"{stem}.csv"
            with path.open("w", newline="", encoding="utf-8-sig") as fh:
                writer = csv.DictWriter(fh, fieldnames=SEARCH_FIELDS, extrasaction="ignore")
                writer.writeheader()
                for row in rows:
                    writer.writerow({f: _stringify(row.get(f)) for f in SEARCH_FIELDS})
        elif fmt == "xlsx":
            path = out / f"{stem}.xlsx"
            _write_search_xlsx(path, rows)
        elif fmt == "md":
            path = out / f"{stem}.md"
            table = _markdown_table(
                [f.replace("_", " ").title() for f in SEARCH_FIELDS],
                [[row.get(f) for f in SEARCH_FIELDS] for row in rows],
            )
            path.write_text(f"# Companies ({len(rows)})\n\n{table}", encoding="utf-8")
        elif fmt == "txt":
            path = out / f"{stem}.txt"
            lines = [f"{len(rows)} companies", "=" * 70]
            for row in rows:
                lines.append(f"{_stringify(row['company_number']):<10} {_stringify(row['company_name'])}")
                lines.append(f"{'':<10} {_stringify(row['company_status'])} | inc. {_stringify(row['date_of_creation'])} | {_stringify(row['registered_address'])}")
            path.write_text("\n".join(lines), encoding="utf-8")
        else:
            raise KYBError(f"Unknown format '{fmt}'. Choose from: {', '.join(SUPPORTED_FORMATS)}, all")
        written.append(path)
    return written


def _write_search_xlsx(path: Path, rows: list) -> None:
    sheet = [("Companies", SEARCH_FIELDS, [[row.get(f) for f in SEARCH_FIELDS] for row in rows])]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        _write_xlsx_stdlib(path, sheet)
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"
    ws.append(SEARCH_FIELDS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_stringify(row.get(f)) for f in SEARCH_FIELDS])
    ws.freeze_panes = "A2"
    for idx, column in enumerate(SEARCH_FIELDS, start=1):
        longest = max([len(column)] + [len(_stringify(r.get(column))) for r in rows] or [0])
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(longest + 2, 12), 60)
    wb.save(path)


def parse_formats(value: str) -> list:
    if value.strip().lower() == "all":
        return list(SUPPORTED_FORMATS)
    formats = [f.strip().lower() for f in value.split(",") if f.strip()]
    unknown = [f for f in formats if f not in SUPPORTED_FORMATS]
    if unknown:
        raise KYBError(f"Unknown format(s): {', '.join(unknown)}. Choose from: {', '.join(SUPPORTED_FORMATS)}, all")
    return formats or ["json"]


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Companies House KYC/KYB lookup with optional UK sanctions screening.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "examples:\n"
            "  python companies_house_kyc.py 00102498\n"
            "  python companies_house_kyc.py 00102498 09659799 --format json,csv,xlsx\n"
            "  python companies_house_kyc.py --search 'monzo bank'\n"
            "  python companies_house_kyc.py 00102498 --format all --outdir reports\n"
            "  python companies_house_kyc.py --advanced-search --sic 62012 --status active --format csv\n"
            "  python companies_house_kyc.py --advanced-search --location london --limit 500 --full\n"
            "  python companies_house_kyc.py --numbers-file clients.csv --format xlsx\n"
            "  python companies_house_kyc.py --bulk-info\n"
        ),
    )
    parser.add_argument("company_numbers", nargs="*", help="One or more company numbers, e.g. 00102498")
    parser.add_argument("--search", metavar="NAME", help="Search by company name instead of number")
    parser.add_argument("--numbers-file", metavar="PATH", help="Read company numbers from a text/CSV file (one per line)")
    parser.add_argument("--bulk-info", action="store_true", help="Explain how to get every UK company, and exit")

    bulk = parser.add_argument_group(
        "bulk listing (advanced search)",
        "Return every company matching a filter, rather than one you name. "
        "At least one filter is required -- there is no 'list all companies' endpoint.",
    )
    bulk.add_argument("--advanced-search", action="store_true", help="List companies matching the filters below")
    bulk.add_argument("--sic", action="append", metavar="CODE", help="SIC code, repeatable (e.g. --sic 62012)")
    bulk.add_argument("--status", action="append", metavar="STATUS", help="active, dissolved, liquidation ... repeatable")
    bulk.add_argument("--type", action="append", metavar="TYPE", dest="company_type", help="ltd, plc, llp ... repeatable")
    bulk.add_argument("--location", metavar="PLACE", help="Registered office town/area, e.g. london")
    bulk.add_argument("--name-includes", metavar="TEXT", help="Company name contains this text")
    bulk.add_argument("--name-excludes", metavar="TEXT", help="Company name does not contain this text")
    bulk.add_argument("--incorporated-from", metavar="YYYY-MM-DD", help="Incorporated on or after this date")
    bulk.add_argument("--incorporated-to", metavar="YYYY-MM-DD", help="Incorporated on or before this date")
    bulk.add_argument("--dissolved-from", metavar="YYYY-MM-DD", help="Dissolved on or after this date")
    bulk.add_argument("--dissolved-to", metavar="YYYY-MM-DD", help="Dissolved on or before this date")
    bulk.add_argument("--limit", type=int, default=1000, help="Maximum companies to return (default: 1000)")
    bulk.add_argument("--full", action="store_true", help="Run the full KYB check on every company found (slow: ~4 requests each)")
    parser.add_argument("--api-key", help="API key (prefer COMPANIES_HOUSE_API_KEY or a .env file)")
    parser.add_argument("--format", default="json", help=f"Comma-separated: {', '.join(SUPPORTED_FORMATS)}, or 'all' (default: json)")
    parser.add_argument("--outdir", default="kyb_output", help="Output directory (default: kyb_output)")
    parser.add_argument("--stem", help="Base filename for the outputs (default: auto, with a timestamp)")
    parser.add_argument("--no-sanctions", action="store_true", help="Skip UK Sanctions List screening")
    parser.add_argument("--sanctions-file", help="Use a local UK Sanctions List XML instead of downloading")
    parser.add_argument("--threshold", type=float, default=0.85, help="Fuzzy match threshold 0-1 (default: 0.85)")
    parser.add_argument("--no-filings", action="store_true", help="Skip the filing history lookup")
    parser.add_argument("--print", dest="print_json", action="store_true", help="Also print the JSON to stdout")
    return parser


def main(argv: Optional[list] = None) -> int:
    args = build_parser().parse_args(argv)

    if args.bulk_info:
        print(bulk_snapshot_info())
        return 0

    if not (args.company_numbers or args.search or args.advanced_search or args.numbers_file):
        build_parser().print_help()
        return 2

    try:
        session = make_session(load_api_key(args.api_key))
        formats = parse_formats(args.format)
        company_numbers = list(args.company_numbers)

        if args.numbers_file:
            company_numbers.extend(read_numbers_file(args.numbers_file))

        if args.advanced_search:
            print("Searching the register ...", file=sys.stderr)
            hits = advanced_search(
                session,
                name_includes=args.name_includes,
                name_excludes=args.name_excludes,
                sic_codes=args.sic,
                company_status=args.status,
                company_type=args.company_type,
                location=args.location,
                incorporated_from=args.incorporated_from,
                incorporated_to=args.incorporated_to,
                dissolved_from=args.dissolved_from,
                dissolved_to=args.dissolved_to,
                limit=args.limit,
            )
            if not hits:
                print("No companies matched those filters.")
                return 1
            print(f"  {len(hits)} companies retrieved.", file=sys.stderr)

            if not args.full:
                # Directory only -- one row per company, no per-company lookups.
                paths = export_search(hits, formats, args.outdir, args.stem)
                print("\nWritten:")
                for path in paths:
                    print(f"  {path}")
                print("\nAdd --full to also pull officers, PSC and sanctions screening for each.")
                return 0

            company_numbers.extend(h.get("company_number") for h in hits if h.get("company_number"))

        if args.search:
            found = search_company(session, args.search, items=10)
            items = found.get("items") or []
            if not items:
                print(f"No companies matched '{args.search}'.")
                return 1
            print(f"Matches for '{args.search}':\n")
            for item in items:
                print(f"  {item.get('company_number'):<10} {item.get('title')}")
                print(f"  {'':<10} {item.get('company_status', '?')} | {item.get('address_snippet', '')}\n")
            if not company_numbers:
                print("Re-run with one of these company numbers to pull the full KYB record.")
                return 0

        sanctions_list: list = []
        if not args.no_sanctions:
            print("Loading UK Sanctions List ...", file=sys.stderr)
            try:
                sanctions_list = load_uksl_names(args.sanctions_file, cache_dir=args.outdir)
                print(f"  {len(sanctions_list)} designated names loaded.", file=sys.stderr)
            except KYBError as exc:
                # Screening failing must not throw away the Companies House data.
                print(f"  ! Sanctions screening skipped: {exc}", file=sys.stderr)

        if len(company_numbers) > 10:
            # Stay inside 600 requests / 5 minutes without relying on 429 retries.
            globals()["MIN_REQUEST_INTERVAL"] = 0.55
            minutes = len(company_numbers) * 4 * 0.55 / 60
            print(f"  Pacing for the rate limit: roughly {minutes:.0f} minute(s) for {len(company_numbers)} companies.", file=sys.stderr)

        results = []
        for index, number in enumerate(company_numbers, start=1):
            print(f"[{index}/{len(company_numbers)}] Fetching {normalise_company_number(number)} ...", file=sys.stderr)
            try:
                result = run_kyb_check(session, number, include_filings=not args.no_filings)
            except KYBError as exc:
                # One bad number must not lose the whole batch.
                print(f"  ! Skipped {number}: {exc}".splitlines()[0], file=sys.stderr)
                continue
            if sanctions_list:
                screen_kyb_result(result, sanctions_list, args.threshold)
            results.append(result)

        if not results:
            print("\nNo companies could be retrieved.", file=sys.stderr)
            return 1

        paths = export(results, formats, args.outdir, args.stem)

        print("\nWritten:")
        for path in paths:
            print(f"  {path}")

        hits = sum(summary_row(r)["sanctions_hit_count"] for r in results)
        if hits:
            print(f"\n  ** {hits} possible sanctions match(es) -- review manually before relying on this. **")

        if args.print_json:
            print(json.dumps(results[0] if len(results) == 1 else results, indent=2, ensure_ascii=False))

        return 0

    except KYBError as exc:
        print(f"\nError: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("\nInterrupted.", file=sys.stderr)
        return 130


if __name__ == "__main__":
    sys.exit(main())
